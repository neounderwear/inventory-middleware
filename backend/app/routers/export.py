import io
import zipfile
from datetime import datetime
from typing import Optional

import pandas as pd
from fastapi import APIRouter, Depends, Form, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import MasterProduct, POItem, PurchaseOrder, EntityStock

router = APIRouter(prefix="/api/export", tags=["Export"])


@router.post("/update-stock/")
def update_stock(
    entitas: str = Form(...),
    selected_brands: Optional[str] = Form(None),
    hide_zero_stock: bool = Form(True),
    db: Session = Depends(get_db),
):
    if entitas not in ["GUDANG", "JAGOAN", "RJM", "7B"]:
        raise HTTPException(status_code=400, detail="Invalid entitas. Must be GUDANG, JAGOAN, RJM, or 7B.")

    target_entity = "GPD" if entitas == "GUDANG" else entitas
    dest_stocks = db.query(EntityStock).filter(EntityStock.entity == target_entity).all()
    stock_map = {s.sku: s.stock for s in dest_stocks}

    stock_records = []
    for sku, stock in stock_map.items():
        stock_records.append({'sku': sku, 'stock': stock})
    
    df_stock = pd.DataFrame(stock_records)
    if df_stock.empty:
        df_stock = pd.DataFrame(columns=['sku', 'stock'])

    products = (
        db.query(MasterProduct).all()
    )

    product_data = []
    for p in products:
        product_data.append({
            "sku": p.sku,
            "product": p.nama_produk,
            "brand": p.brand,
            "variant": p.varian,
        })

    df_products = pd.DataFrame(product_data)
    if df_products.empty:
        df_products = pd.DataFrame(columns=["sku", "product", "brand", "variant"])

    # Merge master products with stock data
    # We only want items that are present in master_products
    required_columns = ["sku", "stock"]
    df_merged = pd.merge(
        df_products, df_stock[required_columns], on="sku", how="inner"
    )

    # After merging, ensure column names are consistent
    df_merged.columns = df_merged.columns.str.lower().str.strip()
    
    if 'brand' in df_merged.columns:
        df_merged['brand'] = df_merged['brand'].astype(str).str.strip()

    df_merged["stock"] = df_merged["stock"].fillna(0).astype(int)


    # Apply whitelist brand filter — only include selected brands
    if selected_brands:
        import json
        try:
            parsed = json.loads(selected_brands)
            if not isinstance(parsed, list):
                parsed = [str(selected_brands)]
        except json.JSONDecodeError:
            parsed = selected_brands.split(',')
            
        allowed = [str(b).strip().lower() for b in parsed if str(b).strip()]
        df_merged = df_merged[df_merged["brand"].str.lower().isin(allowed)]

    if hide_zero_stock:
        df_merged = df_merged[df_merged["stock"] > 0]

    # Rename columns to match reference file format
    df_merged = df_merged.rename(columns={
        "product": "Produk",
        "sku": "SKU",
        "variant": "Varian",
        "stock": "Stok",
        "brand": "Brand",
    })

    # Reorder columns to match reference
    df_merged = df_merged[["Produk", "SKU", "Varian", "Stok", "Brand"]]



    # Sort alphabetically by Brand then Produk for readability
    df_merged = df_merged.sort_values(by=['Brand', 'Produk'], ascending=[True, True])

    # Generate styled Excel file with openpyxl
    wb = Workbook()
    ws = wb.active
    ws.title = "Update Stock"

    current_date = datetime.now().strftime("%d-%m-%Y")

    # --- Global font ---
    arial_bold_14 = Font(name="Arial", size=14, bold=True)
    arial_bold_10 = Font(name="Arial", size=10, bold=True)
    arial_10 = Font(name="Arial", size=10, bold=False)

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    yellow_fill = PatternFill(
        patternType="solid", fgColor="FFFF00"
    )

    # --- Row 1: Title (UPPERCASE, Arial 14, Bold) ---
    ws.merge_cells("A1:E1")
    title_cell = ws["A1"]
    if entitas == "GUDANG":
        title_cell.value = "UPDATE STOCK CV GUDANG PAKAIAN DALAM"
    elif entitas == "JAGOAN":
        title_cell.value = "UPDATE STOCK JAGOAN"
    elif entitas == "RJM":
        title_cell.value = "UPDATE STOCK RJM"
    elif entitas == "7B":
        title_cell.value = "UPDATE STOCK 7B"
    title_cell.font = arial_bold_14
    title_cell.alignment = Alignment(horizontal="center")

    # --- Row 2: Empty ---

    # --- Row 3: Date (Arial 10, Bold) ---
    ws.merge_cells("A3:E3")
    date_cell = ws["A3"]
    date_cell.value = f"Tanggal: {current_date}"
    date_cell.font = arial_bold_10

    # --- Row 4: Empty ---

    # --- Row 5: Headers (Arial 10, Bold, Yellow fill) ---
    headers_list = ["Produk", "SKU", "Varian", "Stok", "Brand"]
    for col_idx, header in enumerate(headers_list, start=1):
        cell = ws.cell(row=5, column=col_idx, value=header)
        cell.font = arial_bold_10
        cell.fill = yellow_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    # --- Auto-filter on header row ---
    ws.auto_filter.ref = f"A5:E{len(df_merged) + 5}"

    # --- Column widths ---
    ws.column_dimensions["A"].width = 45  # Produk
    ws.column_dimensions["B"].width = 20  # SKU
    ws.column_dimensions["C"].width = 20  # Varian
    ws.column_dimensions["D"].width = 10  # Stok
    ws.column_dimensions["E"].width = 20  # Brand

    # --- Data Rows (Row 6 onwards): Arial 10, not bold ---
    for r_idx, row in enumerate(df_merged.itertuples(index=False), start=6):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.font = arial_10
            cell.border = thin_border

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)

    tanggal_hari_ini = datetime.now().strftime("%d-%m-%Y")
    filename = f"UPDATE STOCK {entitas.upper()} {tanggal_hari_ini}.xlsx"

    response_headers = {
        "Content-Disposition": f'attachment; filename="{filename}"'
    }

    return StreamingResponse(
        excel_buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=response_headers,
    )
